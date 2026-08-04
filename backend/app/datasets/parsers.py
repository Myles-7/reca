from __future__ import annotations

import csv
import hashlib
import json
import re
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import (  # type: ignore[import-untyped]
    InvalidFileException,
)

from app.api.errors import ContractError
from app.datasets import limits
from app.models import DatasetColumnType, DatasetFileFormat

_DELIMITERS = ",;\t|"
_CSV_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "cp1252")
_SENSITIVE_NAME = re.compile(
    r"(^|[_\s-])(name|email|phone|mobile|address|id_?card|身份证|姓名|电话|邮箱|住址)($|[_\s-])",
    re.IGNORECASE,
)
_FORMULA_TAG = re.compile(rb"<f(?:\s|>)")


@dataclass(frozen=True)
class WorksheetInfo:
    name: str
    ordinal: int
    visibility: str
    estimated_rows: int
    estimated_columns: int
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class ColumnProfile:
    source_name: str
    inferred_type: DatasetColumnType
    unique_count: int
    missing_ratio: float
    example_values: tuple[Any, ...]
    is_sensitive: bool


@dataclass(frozen=True)
class ParsedTable:
    headers: tuple[str, ...]
    row_count: int
    columns: tuple[ColumnProfile, ...]
    preview_rows: tuple[dict[str, Any], ...]
    warnings: tuple[str, ...]
    projection_hash: str
    schema_hash: str


def _error(code: str, message: str, *, status_code: int = 422) -> ContractError:
    return ContractError(status_code=status_code, code=code, message=message)


def detect_format(
    filename: str, mime_type: str | None, prefix: bytes
) -> DatasetFileFormat:
    suffix = Path(filename).suffix.lower()
    mime = (mime_type or "").split(";", 1)[0].strip().lower()
    if suffix == ".csv":
        if prefix.startswith(b"PK\x03\x04") or mime not in {
            "text/csv",
            "application/csv",
            "text/plain",
            "application/octet-stream",
            "",
        }:
            raise _error(
                "FILE_TYPE_MISMATCH", "The MIME type or content does not match CSV."
            )
        if mime in {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/zip",
        }:
            raise _error("FILE_TYPE_MISMATCH", "The file content does not match CSV.")
        return DatasetFileFormat.CSV
    if suffix == ".xlsx":
        if not prefix.startswith(b"PK\x03\x04"):
            raise _error("FILE_TYPE_MISMATCH", "The file content does not match XLSX.")
        if mime and mime not in {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/octet-stream",
            "application/zip",
        }:
            raise _error("FILE_TYPE_MISMATCH", "The MIME type does not match XLSX.")
        return DatasetFileFormat.XLSX
    raise _error("UNSUPPORTED_FILE_TYPE", "Only CSV and XLSX files are supported.")


def _xlsx_preflight(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if not members or len(members) > limits.MAX_ZIP_MEMBERS:
                raise _error(
                    "WORKBOOK_LIMIT_EXCEEDED",
                    "The workbook ZIP member limit was exceeded.",
                )
            expanded = 0
            for member in members:
                normalized = member.filename.replace("\\", "/")
                if normalized.startswith("/") or "../" in normalized.split("/"):
                    raise _error(
                        "INVALID_WORKBOOK", "The workbook contains an unsafe ZIP path."
                    )
                expanded += member.file_size
                if member.compress_size == 0 and member.file_size > 0:
                    raise _error(
                        "WORKBOOK_COMPRESSION_RISK",
                        "The workbook has an unsafe compression ratio.",
                    )
                if (
                    member.compress_size
                    and member.file_size / member.compress_size > limits.MAX_ZIP_RATIO
                ):
                    raise _error(
                        "WORKBOOK_COMPRESSION_RISK",
                        "The workbook has an unsafe compression ratio.",
                    )
            if expanded > limits.MAX_EXPANDED_XLSX_BYTES:
                raise _error(
                    "WORKBOOK_LIMIT_EXCEEDED", "The expanded workbook is too large."
                )
            names = {member.filename.lower() for member in members}
            if any(name.endswith("vbaproject.bin") for name in names):
                raise _error(
                    "UNSAFE_WORKBOOK", "Macro-enabled workbook content is not allowed."
                )
            if any(name.startswith("xl/externallinks/") for name in names):
                raise _error(
                    "UNSAFE_WORKBOOK", "External workbook links are not allowed."
                )
            for member in members:
                if member.filename.lower().endswith(".rels"):
                    with archive.open(member) as relationships:
                        tail = b""
                        while chunk := relationships.read(64 * 1024):
                            candidate = (tail + chunk).lower()
                            if (
                                b'targetmode="external"' in candidate
                                or b"targetmode='external'" in candidate
                            ):
                                raise _error(
                                    "UNSAFE_WORKBOOK",
                                    "External workbook relationships are not allowed.",
                                )
                            tail = candidate[-24:]
                if member.filename.startswith(
                    "xl/worksheets/"
                ) and member.filename.endswith(".xml"):
                    with archive.open(member) as worksheet:
                        tail = b""
                        while chunk := worksheet.read(64 * 1024):
                            candidate = tail + chunk
                            if _FORMULA_TAG.search(candidate):
                                raise _error(
                                    "UNSAFE_WORKBOOK",
                                    "Workbook formulas are not allowed for import.",
                                )
                            tail = candidate[-2:]
    except zipfile.BadZipFile as exc:
        raise _error("INVALID_WORKBOOK", "The XLSX workbook is damaged.") from exc


def worksheet_manifest(path: Path) -> tuple[WorksheetInfo, ...]:
    _xlsx_preflight(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except (
        OSError,
        ValueError,
        KeyError,
        zipfile.BadZipFile,
        InvalidFileException,
    ) as exc:
        raise _error("INVALID_WORKBOOK", "The XLSX workbook is damaged.") from exc
    try:
        if len(workbook.worksheets) > limits.MAX_WORKSHEETS:
            raise _error(
                "WORKSHEET_LIMIT_EXCEEDED", "The workbook has too many worksheets."
            )
        result = []
        for ordinal, sheet in enumerate(workbook.worksheets, start=1):
            if (
                sheet.max_column > limits.MAX_COLUMNS
                or sheet.max_row - 1 > limits.MAX_ROWS
            ):
                raise _error(
                    "DATASET_LIMIT_EXCEEDED",
                    "A worksheet exceeds the row or column limit.",
                )
            result.append(
                WorksheetInfo(
                    name=sheet.title,
                    ordinal=ordinal,
                    visibility=sheet.sheet_state.upper(),
                    estimated_rows=max(sheet.max_row - 1, 0),
                    estimated_columns=sheet.max_column,
                    warnings=("HIDDEN_WORKSHEET",)
                    if sheet.sheet_state != "visible"
                    else (),
                )
            )
        return tuple(result)
    finally:
        workbook.close()


def _csv_encoding(path: Path) -> str:
    with path.open("rb") as source:
        sample = source.read(65_536)
    for encoding in _CSV_ENCODINGS:
        try:
            sample.decode(encoding)
        except UnicodeDecodeError:
            continue
        return encoding
    raise _error("UNSUPPORTED_ENCODING", "The CSV encoding is not supported.")


def csv_encoding(path: Path) -> str:
    """Return the validated deterministic CSV encoding for downstream readers."""
    return _csv_encoding(path)


def _clean_headers(values: Iterable[Any]) -> tuple[str, ...]:
    headers = tuple(str(value).strip() if value is not None else "" for value in values)
    if not headers or all(not value for value in headers):
        raise _error("EMPTY_DATASET", "The dataset has no header row.")
    if any(not value for value in headers):
        raise _error("INVALID_HEADER", "Every dataset column must have a name.")
    if len(headers) > limits.MAX_COLUMNS:
        raise _error("COLUMN_LIMIT_EXCEEDED", "The dataset has too many columns.")
    if len(set(headers)) != len(headers):
        raise _error("DUPLICATE_COLUMN", "Dataset column names must be unique.")
    return headers


def _value_type(value: Any) -> DatasetColumnType:
    if value is None or value == "":
        return DatasetColumnType.UNKNOWN
    if isinstance(value, bool):
        return DatasetColumnType.BOOLEAN
    if isinstance(value, int):
        return DatasetColumnType.INTEGER
    if isinstance(value, float):
        return DatasetColumnType.NUMERIC
    if isinstance(value, datetime):
        return DatasetColumnType.DATETIME
    if isinstance(value, date):
        return DatasetColumnType.DATE
    text = str(value).strip()
    lowered = text.lower()
    if lowered in {"true", "false", "yes", "no"}:
        return DatasetColumnType.BOOLEAN
    try:
        int(text)
        return DatasetColumnType.INTEGER
    except ValueError:
        pass
    try:
        float(text)
        return DatasetColumnType.NUMERIC
    except ValueError:
        return DatasetColumnType.STRING


def _final_type(
    types: set[DatasetColumnType], unique_count: int, non_missing: int
) -> DatasetColumnType:
    types.discard(DatasetColumnType.UNKNOWN)
    if not types:
        return DatasetColumnType.UNKNOWN
    if types <= {DatasetColumnType.INTEGER}:
        return DatasetColumnType.INTEGER
    if types <= {DatasetColumnType.INTEGER, DatasetColumnType.NUMERIC}:
        return DatasetColumnType.NUMERIC
    if len(types) == 1:
        inferred = next(iter(types))
        if (
            inferred == DatasetColumnType.STRING
            and non_missing
            and unique_count <= min(50, max(2, non_missing // 5))
        ):
            return DatasetColumnType.CATEGORY
        return inferred
    return DatasetColumnType.STRING


def _safe_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime | date):
        return value.isoformat()
    text = str(value)
    if len(text) > limits.MAX_CELL_CHARS:
        raise _error("CELL_LIMIT_EXCEEDED", "A cell exceeds the maximum length.")
    return value if isinstance(value, bool | int | float) else text


def redact_risky_value(value: Any) -> Any:
    if isinstance(value, str) and (
        value[:1] in {"=", "+", "@"}
        or value.lower().startswith(("http://", "https://"))
    ):
        return "[REDACTED]"
    return value


def _profile(
    headers: tuple[str, ...],
    rows: Iterable[Iterable[Any]],
    *,
    projection: str,
    preview_offset: int = 0,
) -> ParsedTable:
    missing = [0] * len(headers)
    non_missing = [0] * len(headers)
    unique: list[set[str]] = [set() for _ in headers]
    examples: list[list[Any]] = [[] for _ in headers]
    types: list[set[DatasetColumnType]] = [set() for _ in headers]
    preview: list[dict[str, Any]] = []
    formula_candidates = 0
    row_count = 0
    for raw_row in rows:
        row = list(raw_row)
        if not row or all(value is None or value == "" for value in row):
            continue
        if len(row) > len(headers):
            raise _error(
                "COLUMN_LIMIT_EXCEEDED", "A data row has more columns than the header."
            )
        row.extend([None] * (len(headers) - len(row)))
        row_count += 1
        if row_count > limits.MAX_ROWS:
            raise _error("ROW_LIMIT_EXCEEDED", "The dataset has too many rows.")
        output: dict[str, Any] = {}
        for index, raw in enumerate(row):
            value = _safe_cell(raw)
            if value is None or value == "":
                missing[index] += 1
            else:
                non_missing[index] += 1
                canonical = json.dumps(
                    value, ensure_ascii=False, sort_keys=True, default=str
                )
                unique[index].add(canonical)
                types[index].add(_value_type(value))
                if (
                    len(examples[index]) < limits.EXAMPLE_VALUE_LIMIT
                    and value not in examples[index]
                ):
                    examples[index].append(value)
                if isinstance(value, str) and value[:1] in {"=", "+", "@"}:
                    formula_candidates += 1
            if row_count > preview_offset and len(preview) < limits.PREVIEW_MAX_ROWS:
                output[headers[index]] = value
        if row_count > preview_offset and len(preview) < limits.PREVIEW_MAX_ROWS:
            preview.append(output)
    if row_count == 0:
        raise _error("EMPTY_DATASET", "The dataset has no data rows.")
    profiles = []
    for index, header in enumerate(headers):
        sensitive = bool(_SENSITIVE_NAME.search(header))
        safe_examples = tuple(
            "***" if sensitive else redact_risky_value(item) for item in examples[index]
        )
        profiles.append(
            ColumnProfile(
                source_name=header,
                inferred_type=_final_type(
                    types[index], len(unique[index]), non_missing[index]
                ),
                unique_count=len(unique[index]),
                missing_ratio=round(missing[index] / row_count, 8),
                example_values=safe_examples,
                is_sensitive=sensitive,
            )
        )
    projection_hash = hashlib.sha256(projection.encode("utf-8")).hexdigest()
    schema_payload = [
        (profile.source_name, profile.inferred_type.value) for profile in profiles
    ]
    schema_hash = hashlib.sha256(
        json.dumps(schema_payload, separators=(",", ":"), ensure_ascii=False).encode()
    ).hexdigest()
    warnings = ("CSV_FORMULA_PREFIX_CANDIDATE",) if formula_candidates else ()
    return ParsedTable(
        headers,
        row_count,
        tuple(profiles),
        tuple(preview),
        warnings,
        projection_hash,
        schema_hash,
    )


def parse_csv(path: Path, *, preview_offset: int = 0) -> ParsedTable:
    encoding = _csv_encoding(path)
    with path.open("r", encoding=encoding, newline="") as source:
        sample = source.read(65_536)
        source.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=_DELIMITERS)
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(source, dialect)
        try:
            headers = _clean_headers(next(reader))
        except StopIteration as exc:
            raise _error("EMPTY_DATASET", "The CSV file is empty.") from exc
        try:
            return _profile(
                headers,
                reader,
                projection=f"csv:{encoding}:{dialect.delimiter}",
                preview_offset=preview_offset,
            )
        except UnicodeDecodeError as exc:
            raise _error(
                "UNSUPPORTED_ENCODING", "The CSV encoding is not supported."
            ) from exc
        except csv.Error as exc:
            raise _error("INVALID_CSV", "The CSV structure is invalid.") from exc


def parse_xlsx(
    path: Path, worksheet_name: str, *, preview_offset: int = 0
) -> ParsedTable:
    _xlsx_preflight(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except (
        OSError,
        ValueError,
        KeyError,
        zipfile.BadZipFile,
        InvalidFileException,
    ) as exc:
        raise _error("INVALID_WORKBOOK", "The XLSX workbook is damaged.") from exc
    try:
        if worksheet_name not in workbook.sheetnames:
            raise _error(
                "WORKSHEET_NOT_FOUND", "The selected worksheet does not exist."
            )
        sheet = workbook[worksheet_name]
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = _clean_headers(next(iterator))
        except StopIteration as exc:
            raise _error("EMPTY_DATASET", "The selected worksheet is empty.") from exc
        return _profile(
            headers,
            iterator,
            projection=f"xlsx:{worksheet_name}",
            preview_offset=preview_offset,
        )
    finally:
        workbook.close()
