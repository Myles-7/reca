from __future__ import annotations

import hashlib
import re
from io import BytesIO
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document

from app.models import ManuscriptIssue, ManuscriptIssueType

FIXER_VERSION = "m6-low-risk-1.0"
ALLOWED_FIXERS = {
    ManuscriptIssueType.PUNCTUATION_ISSUE: "NORMALIZE_SPACES",
    ManuscriptIssueType.UNIT_FORMAT_ISSUE: "NORMALIZE_UNIT_SPACING",
}


def canonical_actions(issues: list[ManuscriptIssue]) -> list[dict[str, Any]]:
    actions = []
    for issue in sorted(issues, key=lambda item: str(item.id)):
        fixer = ALLOWED_FIXERS.get(issue.issue_type)
        paragraph = issue.locator.get("paragraph")
        if fixer is None or not issue.auto_fixable or not isinstance(paragraph, int):
            raise ValueError("Issue is not eligible for an allowlisted automatic fix.")
        actions.append(
            {
                "issue_id": str(issue.id),
                "fixer": fixer,
                "paragraph": paragraph,
                "finding_hash": issue.finding_hash,
            }
        )
    return actions


def _replace(value: str, fixer: str) -> str:
    if fixer == "NORMALIZE_SPACES":
        return re.sub(r" {2,}", " ", value)
    if fixer == "NORMALIZE_UNIT_SPACING":
        return re.sub(r"(?<=\d)(?=(?:kg|mg|cm|mm|%)(?:\b|$))", " ", value)
    raise ValueError("Unknown fixer.")


def _protected_parts(content: bytes) -> dict[str, str]:
    with ZipFile(BytesIO(content)) as package:
        return {
            name: hashlib.sha256(package.read(name)).hexdigest()
            for name in package.namelist()
            if not name.startswith("word/") and name != "[Content_Types].xml"
        }


def _deterministic_package(content: bytes) -> bytes:
    output = BytesIO()
    with (
        ZipFile(BytesIO(content)) as source,
        ZipFile(output, "w", compression=ZIP_DEFLATED) as target,
    ):
        for name in sorted(source.namelist()):
            original = source.getinfo(name)
            item = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            item.compress_type = ZIP_DEFLATED
            item.external_attr = original.external_attr
            item.create_system = original.create_system
            target.writestr(item, source.read(name))
    return output.getvalue()


def apply_fixes(
    content: bytes,
    actions: list[dict[str, Any]],
    *,
    unsupported_features: list[str],
    unknown_parts: list[str] | None = None,
) -> tuple[bytes, list[dict[str, Any]]]:
    if unsupported_features or unknown_parts:
        raise ValueError("Documents with unsupported structures are not auto-fixable.")
    document = Document(BytesIO(content))
    changes: list[dict[str, Any]] = []
    for action in actions:
        index = int(action["paragraph"])
        if index < 0 or index >= len(document.paragraphs):
            raise ValueError("Fix locator is stale.")
        paragraph = document.paragraphs[index]
        before = paragraph.text
        for run in paragraph.runs:
            run.text = _replace(run.text, str(action["fixer"]))
        after = paragraph.text
        if before == after:
            raise ValueError("Fix locator no longer matches fixable content.")
        changes.append(
            {
                "issue_id": action["issue_id"],
                "paragraph": index,
                "before": before[:240],
                "after": after[:240],
                "fixer": action["fixer"],
            }
        )
    output = BytesIO()
    document.save(output)
    result = _deterministic_package(output.getvalue())
    if _protected_parts(content) != _protected_parts(result):
        raise ValueError("Protected package parts were not preserved.")
    return result, changes


def preview_fixes(
    content: bytes,
    actions: list[dict[str, Any]],
    *,
    unsupported_features: list[str],
    unknown_parts: list[str] | None = None,
) -> dict[str, Any]:
    output, changes = apply_fixes(
        content,
        actions,
        unsupported_features=unsupported_features,
        unknown_parts=unknown_parts,
    )
    return {
        "schema": "reca.manuscript.fix-preview.v1",
        "fixer_version": FIXER_VERSION,
        "changes": changes,
        "output_sha256": hashlib.sha256(output).hexdigest(),
    }
