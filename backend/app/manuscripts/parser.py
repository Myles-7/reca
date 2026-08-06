from __future__ import annotations

import hashlib
import multiprocessing as mp
import re
from pathlib import Path
from typing import Any, cast
from zipfile import ZipFile

from defusedxml import ElementTree
from docx import Document

PARSER_VERSION = "reca-docx-1.0"
PARSER_TIMEOUT_SECONDS = 60

_BIBLIOGRAPHY_NAMESPACE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/bibliography"
)
_CUSTOM_XML_NAMESPACE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/customXml"
)
_RELATIONSHIPS_NAMESPACE = (
    "http://schemas.openxmlformats.org/package/2006/relationships"
)
_BIBLIOGRAPHY_PARTS = {
    "customXml/item1.xml",
    "customXml/itemProps1.xml",
    "customXml/_rels/item1.xml.rels",
}


def _validated_snapshot(value: Any) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise ValueError("DOCX parser returned an invalid snapshot payload.")
    return cast(dict[str, Any], value)


def _trusted_bibliography_parts(package: ZipFile) -> set[str]:
    if not _BIBLIOGRAPHY_PARTS.issubset(package.namelist()):
        return set()
    try:
        sources = ElementTree.fromstring(package.read("customXml/item1.xml"))
        properties = ElementTree.fromstring(package.read("customXml/itemProps1.xml"))
        relationships = ElementTree.fromstring(
            package.read("customXml/_rels/item1.xml.rels")
        )
    except ElementTree.ParseError, ValueError:
        return set()
    schema_refs = properties.findall(f".//{{{_CUSTOM_XML_NAMESPACE}}}schemaRef")
    relationship_rows = relationships.findall(
        f"{{{_RELATIONSHIPS_NAMESPACE}}}Relationship"
    )
    if (
        sources.tag != f"{{{_BIBLIOGRAPHY_NAMESPACE}}}Sources"
        or not any(
            row.get(f"{{{_CUSTOM_XML_NAMESPACE}}}uri") == _BIBLIOGRAPHY_NAMESPACE
            for row in schema_refs
        )
        or len(relationship_rows) != 1
        or relationship_rows[0].get("Type")
        != "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps"
        or relationship_rows[0].get("Target") != "itemProps1.xml"
    ):
        return set()
    return set(_BIBLIOGRAPHY_PARTS)


def _locator(
    *, paragraph: int | None = None, table: int | None = None, cell: str | None = None
) -> dict[str, Any]:
    return {
        "schema": "reca.manuscript.locator.v1",
        "paragraph": paragraph,
        "table": table,
        "cell": cell,
    }


def _parse_docx_unbounded(path: Path) -> dict[str, Any]:
    document = Document(str(path))
    paragraphs: list[dict[str, Any]] = [
        {
            "index": index,
            "text": paragraph.text,
            "style": paragraph.style.name if paragraph.style else None,
            "locator": _locator(paragraph=index),
        }
        for index, paragraph in enumerate(document.paragraphs)
    ]
    tables: list[dict[str, Any]] = []
    for table_index, table in enumerate(document.tables):
        rows = []
        for row_index, row in enumerate(table.rows):
            rows.append(
                [
                    {
                        "text": cell.text,
                        "locator": _locator(
                            table=table_index, cell=f"R{row_index}C{column_index}"
                        ),
                    }
                    for column_index, cell in enumerate(row.cells)
                ]
            )
        tables.append({"index": table_index, "rows": rows})

    unsupported: list[str] = []
    with ZipFile(path) as package:
        document_xml = package.read("word/document.xml")
        markers = {
            b"<w:ins": "TRACKED_INSERTION",
            b"<w:del": "TRACKED_DELETION",
            b"<w:fldSimple": "FIELD",
            b"<w:instrText": "COMPLEX_FIELD",
            b"<w:txbxContent": "TEXT_BOX",
        }
        unsupported.extend(
            code for marker, code in markers.items() if marker in document_xml
        )
        uses_numbering_style = any(
            paragraph.style is not None
            and paragraph.style.style_id.lower().startswith(
                ("listnumber", "listbullet")
            )
            for paragraph in document.paragraphs
        )
        if b"<w:numPr" in document_xml or uses_numbering_style:
            unsupported.append("NUMBERING")
        names = set(package.namelist())
        trusted_bibliography_parts = _trusted_bibliography_parts(package)
        unknown_parts = sorted(
            name
            for name in names
            if not name.startswith(("word/", "_rels/", "docProps/"))
            and name != "[Content_Types].xml"
            and name not in trusted_bibliography_parts
        )
        advanced_parts = sorted(
            name
            for name in names
            if name.startswith(
                (
                    "word/header",
                    "word/footer",
                    "word/comments",
                    "word/footnotes",
                    "word/endnotes",
                    "word/glossary",
                )
            )
        )
        unsupported.extend(f"UNPROVEN_PART:{name}" for name in advanced_parts)
        unknown_parts.extend(advanced_parts)
    canonical = "\n".join(str(item["text"]) for item in paragraphs).encode("utf-8")
    return {
        "schema": "reca.manuscript.parse.v1",
        "parser_version": PARSER_VERSION,
        "paragraphs": paragraphs,
        "tables": tables,
        "unsupported_features": sorted(set(unsupported)),
        "unknown_parts": unknown_parts,
        "confidence": "LOW" if unsupported else "HIGH",
        "text_hash": hashlib.sha256(canonical).hexdigest(),
    }


def _parse_worker(path_string: str, queue: Any) -> None:
    try:
        queue.put((True, _parse_docx_unbounded(Path(path_string))))
    except BaseException as exc:  # worker boundary must serialize failures
        queue.put((False, f"{type(exc).__name__}: {exc}"))


def parse_docx(
    path: Path, *, timeout_seconds: int = PARSER_TIMEOUT_SECONDS
) -> dict[str, Any]:
    """Parse in an isolated process with a hard deadline.

    The worker is terminated on timeout so a malformed but package-limit-compliant
    OOXML document cannot pin the shared Celery process indefinitely.
    """
    context = mp.get_context("spawn")
    queue = context.Queue(maxsize=1)
    process = context.Process(target=_parse_worker, args=(str(path), queue))
    process.start()
    process.join(timeout_seconds)
    if process.is_alive():
        process.terminate()
        process.join(5)
        raise TimeoutError(f"DOCX parsing exceeded {timeout_seconds} seconds.")
    if queue.empty():
        raise RuntimeError("DOCX parser worker exited without a result.")
    ok, value = queue.get()
    if not ok:
        raise ValueError(f"DOCX parsing failed: {value}")
    return _validated_snapshot(value)


def normalized_text(snapshot: dict[str, Any]) -> str:
    return "\n".join(
        str(item.get("text", "")) for item in snapshot.get("paragraphs", [])
    )


def excerpt(value: str, *, limit: int = 240) -> str:
    return re.sub(r"\s+", " ", value).strip()[:limit]
