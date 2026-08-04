from __future__ import annotations

import csv
import re
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd  # type: ignore[import-untyped]
import pandera.pandas as pa
from openpyxl import load_workbook  # type: ignore[import-untyped]
from pandera.errors import SchemaErrors

from app.data_quality.registry import RuleDefinition, RuleSet
from app.datasets import limits, parsers
from app.models import (
    DataQualityIssueType,
    DataQualitySeverity,
    DatasetColumn,
    DatasetColumnType,
    DatasetFileFormat,
    DatasetSemanticRole,
    DatasetVersion,
)

MAX_AFFECTED_ROWS = 100
MAX_EXAMPLES = 5
MAX_EXAMPLE_CHARS = 64

_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_PHONE = re.compile(r"^\+?[0-9][0-9\s()-]{6,18}$")
_ID_CARD = re.compile(r"^(?:\d{15}|\d{17}[0-9Xx])$")
_STUDENT_NUMBER = re.compile(r"^[A-Za-z]?[0-9]{6,14}$")
_DATEISH = re.compile(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}(?:[ T].*)?$")


@dataclass(frozen=True)
class ColumnContext:
    id: uuid.UUID | None
    source_name: str
    column_order: int
    inferred_type: DatasetColumnType
    confirmed_type: DatasetColumnType | None = None
    semantic_role: DatasetSemanticRole | None = None
    unit: str | None = None
    missing_codes: tuple[str, ...] = ()
    is_identifier: bool = False
    is_sensitive: bool = False

    @classmethod
    def from_model(cls, column: DatasetColumn) -> ColumnContext:
        return cls(
            id=column.id,
            source_name=column.source_name,
            column_order=column.column_order,
            inferred_type=column.inferred_type,
            confirmed_type=column.confirmed_type,
            semantic_role=column.semantic_role,
            unit=column.unit,
            missing_codes=tuple(column.missing_codes or ()),
            is_identifier=column.is_identifier,
            is_sensitive=column.is_sensitive,
        )


@dataclass(frozen=True)
class IssueFact:
    rule_code: str
    issue_type: DataQualityIssueType
    severity: DataQualitySeverity
    column_id: uuid.UUID | None
    column_name: str | None
    affected_row_count: int
    affected_rows: tuple[Any, ...]
    evidence: dict[str, Any]
    description: str
    requires_approval: bool = False


def _pandera_failure_indices(mask: pd.Series, *, rule_code: str) -> tuple[Any, ...]:
    check_frame = pd.DataFrame(
        {"valid": (~mask.astype(bool)).to_numpy(copy=True)},
        index=mask.index.copy(),
    )
    schema = pa.DataFrameSchema(
        {
            "valid": pa.Column(
                bool,
                checks=pa.Check.equal_to(True, error=rule_code),
                nullable=False,
                coerce=False,
            )
        },
        coerce=False,
        strict=True,
    )
    try:
        schema.validate(check_frame, lazy=True, inplace=False)
        return ()
    except SchemaErrors as error:
        cases = error.failure_cases
        indexes = [value for value in cases["index"].tolist() if pd.notna(value)]
        return tuple(dict.fromkeys(indexes))


def _missing_mask(series: pd.Series, codes: tuple[str, ...]) -> pd.Series:
    normalized = series.astype("string").str.strip().str.upper()
    return series.isna() | normalized.isin({code.strip().upper() for code in codes})


def _masked_example(value: Any, *, sensitive: bool) -> Any:
    if sensitive:
        return "[MASKED]"
    if value is None or pd.isna(value):
        return None
    redacted = parsers.redact_risky_value(value)
    text = str(redacted)
    return text[:MAX_EXAMPLE_CHARS]


def _issue(
    *,
    rule: RuleDefinition,
    ruleset: RuleSet,
    mask: pd.Series,
    frame: pd.DataFrame,
    column: ColumnContext | None,
    description: str,
    extra_evidence: dict[str, Any] | None = None,
    sensitive: bool = False,
) -> IssueFact | None:
    failure_indexes = _pandera_failure_indices(mask, rule_code=rule.rule_id)
    if not failure_indexes:
        return None
    affected_count = len(failure_indexes)
    bounded_rows = failure_indexes[:MAX_AFFECTED_ROWS]
    examples: list[Any] = []
    if column is not None and column.source_name in frame.columns:
        for index in failure_indexes:
            value = frame.at[index, column.source_name]
            example = _masked_example(value, sensitive=sensitive or column.is_sensitive)
            if example not in examples:
                examples.append(example)
            if len(examples) == MAX_EXAMPLES:
                break
    evidence = {
        "schema_version": "1.0",
        "ruleset_id": ruleset.ruleset_id,
        "ruleset_version": ruleset.version,
        "ruleset_hash": ruleset.content_hash,
        "rule_id": rule.rule_id,
        "parameters": rule.parameters,
        "column_name": column.source_name if column else None,
        "row_locator": list(bounded_rows),
        "examples": examples,
        "affected_rows_truncated": affected_count > len(bounded_rows),
        "examples_truncated": affected_count > len(examples),
        "pandera": {"lazy": True, "coerce": False},
        **(extra_evidence or {}),
    }
    return IssueFact(
        rule_code=rule.rule_id,
        issue_type=rule.issue_type,
        severity=rule.severity,
        column_id=column.id if column else None,
        column_name=column.source_name if column else None,
        affected_row_count=affected_count,
        affected_rows=tuple(bounded_rows),
        evidence=evidence,
        description=description,
    )


def _token_family(value: Any) -> str:
    if value is None or pd.isna(value) or str(value).strip() == "":
        return "missing"
    text = str(value).strip()
    if text.casefold() in {"true", "false", "yes", "no"}:
        return "boolean"
    try:
        float(text)
        return "numeric"
    except ValueError:
        pass
    if _DATEISH.fullmatch(text):
        return "date"
    return "text"


def _matches_sensitive_pattern(value: Any) -> bool:
    if pd.isna(value) or not str(value).strip():
        return False
    text = str(value).strip()
    return bool(
        _EMAIL.fullmatch(text)
        or (_PHONE.fullmatch(text) and not _DATEISH.fullmatch(text))
        or _ID_CARD.fullmatch(text)
        or _STUDENT_NUMBER.fullmatch(text)
    )


def _numeric(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce")


def _is_date_column(column: ColumnContext) -> bool:
    declared = column.confirmed_type or column.inferred_type
    if declared in {DatasetColumnType.DATE, DatasetColumnType.DATETIME}:
        return True
    lowered = column.source_name.casefold()
    return any(token in lowered for token in ("date", "time", "日期", "时间"))


def scan_dataframe(
    frame: pd.DataFrame,
    *,
    columns: tuple[ColumnContext, ...],
    ruleset: RuleSet,
    include_sensitive_field_detection: bool,
) -> tuple[IssueFact, ...]:
    source_snapshot = frame.copy(deep=True)
    working = frame.copy(deep=True)
    contexts = {column.source_name: column for column in columns}
    facts: list[IssueFact] = []

    for rule in ruleset.rules:
        if rule.issue_type == DataQualityIssueType.MISSING_VALUE:
            defaults = tuple(rule.parameters["default_missing_codes"])
            for column in columns:
                codes = tuple(dict.fromkeys((*defaults, *column.missing_codes)))
                mask = _missing_mask(working[column.source_name], codes)
                fact = _issue(
                    rule=rule,
                    ruleset=ruleset,
                    mask=mask,
                    frame=working,
                    column=column,
                    description="Missing or declared missing-code values were observed and require review.",
                    extra_evidence={"missing_codes_checked": list(codes)},
                    sensitive=column.is_sensitive,
                )
                if fact:
                    facts.append(fact)
        elif rule.issue_type == DataQualityIssueType.DUPLICATE_ROW:
            mask = working.duplicated(keep=False)
            fact = _issue(
                rule=rule,
                ruleset=ruleset,
                mask=mask,
                frame=working,
                column=None,
                description="Duplicate rows were observed and require review.",
            )
            if fact:
                facts.append(fact)
        elif rule.issue_type == DataQualityIssueType.DUPLICATE_ID:
            for column in columns:
                if not (
                    column.is_identifier
                    or column.semantic_role == DatasetSemanticRole.ID
                ):
                    continue
                series = working[column.source_name]
                missing = _missing_mask(series, column.missing_codes)
                mask = series.duplicated(keep=False) & ~missing
                fact = _issue(
                    rule=rule,
                    ruleset=ruleset,
                    mask=mask,
                    frame=working,
                    column=column,
                    description="Duplicate values violate the confirmed identifier contract.",
                    extra_evidence={"identifier_basis": "confirmed"},
                    sensitive=True,
                )
                if fact:
                    facts.append(fact)
        elif rule.issue_type == DataQualityIssueType.CONSTANT_COLUMN:
            for column in columns:
                series = working[column.source_name]
                non_missing = series[~_missing_mask(series, column.missing_codes)]
                if non_missing.empty or non_missing.nunique(dropna=True) > 1:
                    continue
                mask = ~_missing_mask(series, column.missing_codes)
                fact = _issue(
                    rule=rule,
                    ruleset=ruleset,
                    mask=mask,
                    frame=working,
                    column=column,
                    description="A constant column is a data-quality clue or analysis limitation, not a confirmed error.",
                )
                if fact:
                    facts.append(fact)
        elif rule.issue_type == DataQualityIssueType.MIXED_TYPE:
            for column in columns:
                series = working[column.source_name]
                families = series.map(_token_family)
                present = set(families) - {"missing"}
                if len(present) <= 1 or present <= {"numeric"}:
                    continue
                dominant = families[families != "missing"].value_counts().index[0]
                mask = ~families.isin({"missing", dominant})
                fact = _issue(
                    rule=rule,
                    ruleset=ruleset,
                    mask=mask,
                    frame=working,
                    column=column,
                    description="Incompatible value families were observed in one column.",
                    extra_evidence={
                        "families": sorted(present),
                        "dominant_family": dominant,
                    },
                )
                if fact:
                    facts.append(fact)
        elif rule.issue_type == DataQualityIssueType.CATEGORY_INCONSISTENCY:
            for column in columns:
                series = working[column.source_name].astype("string")
                non_missing = series[~_missing_mask(series, column.missing_codes)]
                if non_missing.empty or non_missing.nunique() > int(
                    rule.parameters["max_cardinality"]
                ):
                    continue
                normalized = non_missing.str.strip().str.casefold()
                variants = (
                    pd.DataFrame({"raw": non_missing, "normalized": normalized})
                    .groupby("normalized", dropna=False)["raw"]
                    .nunique()
                )
                inconsistent = set(variants[variants > 1].index)
                if not inconsistent:
                    continue
                full_normalized = series.str.strip().str.casefold()
                mask = full_normalized.isin(inconsistent)
                fact = _issue(
                    rule=rule,
                    ruleset=ruleset,
                    mask=mask,
                    frame=working,
                    column=column,
                    description="Category spelling or casing variants were observed and require review.",
                    extra_evidence={"normalized_category_count": len(inconsistent)},
                )
                if fact:
                    facts.append(fact)
        elif rule.issue_type == DataQualityIssueType.OUT_OF_RANGE:
            ranges = rule.parameters["ranges"]
            for column in columns:
                lowered = column.source_name.casefold()
                matched = next(
                    (bounds for token, bounds in ranges.items() if token in lowered),
                    None,
                )
                if matched is None:
                    continue
                numeric = _numeric(working[column.source_name])
                lower, upper = float(matched[0]), float(matched[1])
                mask = numeric.notna() & ((numeric < lower) | (numeric > upper))
                fact = _issue(
                    rule=rule,
                    ruleset=ruleset,
                    mask=mask,
                    frame=working,
                    column=column,
                    description="Values outside a name-based expected range were observed as review candidates.",
                    extra_evidence={"lower": lower, "upper": upper, "heuristic": True},
                )
                if fact:
                    facts.append(fact)
        elif rule.issue_type == DataQualityIssueType.EXTREME_VALUE:
            for column in columns:
                numeric = _numeric(working[column.source_name])
                valid = numeric.dropna()
                if len(valid) < int(rule.parameters["minimum_values"]):
                    continue
                q1, q3 = valid.quantile([0.25, 0.75]).tolist()
                iqr = q3 - q1
                if not np.isfinite(iqr) or iqr <= 0:
                    continue
                multiplier = float(rule.parameters["iqr_multiplier"])
                lower, upper = q1 - multiplier * iqr, q3 + multiplier * iqr
                mask = numeric.notna() & ((numeric < lower) | (numeric > upper))
                fact = _issue(
                    rule=rule,
                    ruleset=ruleset,
                    mask=mask,
                    frame=working,
                    column=column,
                    description="IQR identified extreme-value clues; these are not confirmed errors.",
                    extra_evidence={"method": "IQR", "lower": lower, "upper": upper},
                )
                if fact:
                    facts.append(fact)
        elif rule.issue_type == DataQualityIssueType.GROUP_IMBALANCE:
            for column in columns:
                lowered = column.source_name.casefold()
                if not (
                    column.semantic_role == DatasetSemanticRole.GROUP_VARIABLE
                    or any(
                        token in lowered for token in ("group", "gender", "组", "性别")
                    )
                ):
                    continue
                series = working[column.source_name]
                valid = series[~_missing_mask(series, column.missing_codes)]
                minimum = int(rule.parameters["minimum_values"])
                if len(valid) < minimum or valid.nunique() < 2:
                    continue
                counts = valid.value_counts()
                ratio = float(counts.iloc[0] / len(valid))
                if ratio < float(rule.parameters["dominant_ratio"]):
                    continue
                mask = series == counts.index[0]
                fact = _issue(
                    rule=rule,
                    ruleset=ruleset,
                    mask=mask,
                    frame=working,
                    column=column,
                    description="Group imbalance is an analysis limitation clue, not a confirmed data error.",
                    extra_evidence={
                        "dominant_ratio": ratio,
                        "group_count": int(valid.nunique()),
                    },
                )
                if fact:
                    facts.append(fact)
        elif rule.issue_type == DataQualityIssueType.SUSPICIOUS_UNIT:
            suffix_units = rule.parameters["suffix_units"]
            for column in columns:
                lowered = column.source_name.casefold()
                expected = next(
                    (
                        unit
                        for suffix, unit in suffix_units.items()
                        if lowered.endswith(f"_{suffix}") or lowered == suffix
                    ),
                    None,
                )
                if (
                    expected is None
                    or (column.unit or "").casefold() == expected.casefold()
                ):
                    continue
                series = working[column.source_name]
                mask = ~_missing_mask(series, column.missing_codes)
                fact = _issue(
                    rule=rule,
                    ruleset=ruleset,
                    mask=mask,
                    frame=working,
                    column=column,
                    description="The declared or missing unit is inconsistent with the column name and requires review.",
                    extra_evidence={
                        "expected_unit": expected,
                        "declared_unit": column.unit,
                    },
                )
                if fact:
                    facts.append(fact)
        elif rule.issue_type == DataQualityIssueType.INVALID_DATE:
            for column in columns:
                if not _is_date_column(column):
                    continue
                series = working[column.source_name]
                missing = _missing_mask(series, column.missing_codes)
                parsed = pd.to_datetime(series.where(~missing), errors="coerce")
                mask = ~missing & parsed.isna()
                fact = _issue(
                    rule=rule,
                    ruleset=ruleset,
                    mask=mask,
                    frame=working,
                    column=column,
                    description="Values incompatible with the date contract were observed.",
                )
                if fact:
                    facts.append(fact)
        elif (
            rule.issue_type == DataQualityIssueType.POSSIBLE_SENSITIVE_FIELD
            and include_sensitive_field_detection
        ):
            name_tokens = tuple(rule.parameters["name_tokens"])
            for column in columns:
                series = working[column.source_name].astype("string")
                lowered = column.source_name.casefold()
                name_match = any(token.casefold() in lowered for token in name_tokens)
                value_match = series.map(_matches_sensitive_pattern)
                mask = value_match | (
                    pd.Series(True, index=series.index)
                    if name_match
                    else pd.Series(False, index=series.index)
                )
                mask &= ~_missing_mask(series, column.missing_codes)
                fact = _issue(
                    rule=rule,
                    ruleset=ruleset,
                    mask=mask,
                    frame=working,
                    column=column,
                    description="Field name or bounded pattern evidence indicates a possible sensitive field requiring review.",
                    extra_evidence={
                        "name_match": name_match,
                        "pattern_match_count": int(value_match.sum()),
                    },
                    sensitive=True,
                )
                if fact:
                    facts.append(fact)

    pd.testing.assert_frame_equal(
        source_snapshot, frame, check_dtype=True, check_exact=True
    )
    order = {rule.rule_id: index for index, rule in enumerate(ruleset.rules)}
    return tuple(
        sorted(
            facts,
            key=lambda fact: (
                order[fact.rule_code],
                contexts[fact.column_name].column_order if fact.column_name else -1,
                str(fact.column_id or ""),
            ),
        )
    )


def load_dataframe(path: Path, *, version: DatasetVersion) -> pd.DataFrame:
    if version.file_format == DatasetFileFormat.CSV:
        parsed = parsers.parse_csv(path)
        encoding = parsers.csv_encoding(path)
        with path.open("r", encoding=encoding, newline="") as source:
            sample = source.read(65_536)
            source.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            rows = list(csv.reader(source, dialect))
        headers = list(parsed.headers)
        values = []
        for row in rows[1:]:
            if not row or not any(value != "" for value in row):
                continue
            if len(row) > len(headers):
                raise ValueError("CSV row exceeds the validated header width.")
            values.append([*row, *([None] * (len(headers) - len(row)))])
        return pd.DataFrame(values, columns=headers, dtype=object)

    parsers.worksheet_manifest(path)
    worksheet = version.selected_worksheet_name
    if not worksheet:
        raise ValueError("XLSX DatasetVersion has no selected worksheet.")
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook[worksheet]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return pd.DataFrame()
    headers = [str(value) for value in rows[0]]
    values = [
        list(row) for row in rows[1:] if any(value not in {None, ""} for value in row)
    ]
    if len(values) > limits.MAX_ROWS:
        raise ValueError("Dataset exceeds the quality scan row limit.")
    return pd.DataFrame(values, columns=headers, dtype=object)
