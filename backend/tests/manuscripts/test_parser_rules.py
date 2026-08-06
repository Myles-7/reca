from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document

from app.manuscripts.parser import _validated_snapshot, parse_docx
from app.manuscripts.rules import run_rules
from app.models import ManuscriptIssueType

pytestmark = pytest.mark.no_database


def _document(path: Path) -> None:
    document = Document()
    document.add_heading("Results", level=1)
    document.add_paragraph(
        "The treatment causes improvement in all participants (Smith, 2020).  N = 42."
    )
    document.add_heading("References", level=1)
    document.add_paragraph("Jones, A. (2021). Unused. doi:bad-doi")
    document.add_paragraph("Jones, A. (2021). Duplicate.")
    document.save(path)


def test_parser_emits_stable_locators_and_hash(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    _document(path)
    first = parse_docx(path)
    second = parse_docx(path)
    assert first["schema"] == "reca.manuscript.parse.v1"
    assert first["paragraphs"][1]["locator"] == {
        "schema": "reca.manuscript.locator.v1",
        "paragraph": 1,
        "table": None,
        "cell": None,
    }
    assert first["text_hash"] == second["text_hash"]
    assert "UNPROVEN_PART:word/numbering.xml" not in first["unsupported_features"]
    assert first["unknown_parts"] == []


def test_parser_rejects_non_object_worker_payload() -> None:
    with pytest.raises(ValueError, match="invalid snapshot payload"):
        _validated_snapshot([{"schema": "unexpected-list"}])


def test_parser_marks_numbering_only_when_document_uses_it(tmp_path: Path) -> None:
    path = tmp_path / "numbered.docx"
    document = Document()
    document.add_paragraph("Numbered item", style="List Number")
    document.save(path)

    snapshot = parse_docx(path)

    assert "NUMBERING" in snapshot["unsupported_features"]
    assert snapshot["confidence"] == "LOW"


def test_parser_keeps_unrecognized_custom_xml_fail_closed(tmp_path: Path) -> None:
    path = tmp_path / "custom.docx"
    _document(path)
    with ZipFile(path, "a") as package:
        package.writestr("customXml/untrusted.xml", "<untrusted />")

    snapshot = parse_docx(path)

    assert "customXml/untrusted.xml" in snapshot["unknown_parts"]


def test_p0_rules_find_citation_causality_duplicate_doi_and_format(
    tmp_path: Path,
) -> None:
    path = tmp_path / "sample.docx"
    _document(path)
    findings = run_rules(parse_docx(path), {"CITATION", "CAUSALITY", "BASIC_FORMAT"})
    codes = {item.issue_type for item in findings}
    assert ManuscriptIssueType.IN_TEXT_CITATION_MISSING_REFERENCE in codes
    assert ManuscriptIssueType.UNUSED_REFERENCE in codes
    assert ManuscriptIssueType.DUPLICATE_REFERENCE in codes
    assert ManuscriptIssueType.INVALID_DOI_FORMAT in codes
    assert ManuscriptIssueType.CAUSAL_OVERCLAIM in codes
    assert ManuscriptIssueType.POPULATION_OVERGENERALIZATION in codes
    assert ManuscriptIssueType.PUNCTUATION_ISSUE in codes
    assert all(
        not item.auto_fixable for item in findings if item.severity.value == "HIGH"
    )


def test_parser_reports_tracked_changes_as_low_confidence(tmp_path: Path) -> None:
    path = tmp_path / "tracked.docx"
    _document(path)
    import zipfile

    with zipfile.ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
        members["word/document.xml"] = members["word/document.xml"].replace(
            b"<w:t>Results</w:t>", b"<w:ins><w:r><w:t>Inserted</w:t></w:r></w:ins>"
        )
    with zipfile.ZipFile(path, "w") as archive:
        for name, content in members.items():
            archive.writestr(name, content)
    snapshot = parse_docx(path)
    assert snapshot["confidence"] == "LOW"
    assert "TRACKED_INSERTION" in snapshot["unsupported_features"]
