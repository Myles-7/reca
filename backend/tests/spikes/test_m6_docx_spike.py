from __future__ import annotations

import base64
import re
import stat
import xml.etree.ElementTree as StdElementTree
import zipfile
from io import BytesIO
from pathlib import Path, PurePosixPath

import docx
import lxml
import pytest
from defusedxml import ElementTree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

pytestmark = pytest.mark.no_database

DOCX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
)
RELATIONSHIPS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl2nGQAAAAASUVORK5CYII="
)

AUTHOR_YEAR = re.compile(
    r"(?:[（(]\s*)?(?P<author>[A-Za-z][A-Za-z' -]+|[\u4e00-\u9fff]{1,8}?)"
    r"(?:\s+et\s+al\.|等)?\s*[,，]?\s*[（(]?\s*"
    r"(?P<year>(?:19|20)\d{2}[a-z]?)\s*[）)]?",
    re.IGNORECASE,
)
NUMERIC_FACT = re.compile(
    r"(?P<kind>N|n|p|r|β|beta)\s*(?:=|<|>|≤|≥)\s*"
    r"(?P<value>(?:\.\d+|[-+]?\d+(?:\.\d+)?))",
    re.IGNORECASE,
)
FIGURE_REFERENCE = re.compile(r"(?:Figure|Fig\.|图)\s*(?P<number>\d+)", re.IGNORECASE)


def _rewrite_zip(
    path: Path, replacements: dict[str, bytes], additions: dict[str, bytes]
) -> None:
    source = path.read_bytes()
    with zipfile.ZipFile(BytesIO(source)) as original:
        members = {item.filename: original.read(item) for item in original.infolist()}
    members.update(replacements)
    members.update(additions)
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as output:
        for name, value in members.items():
            output.writestr(name, value)


def _add_spike_ooxml(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        document_xml = ElementTree.fromstring(archive.read("word/document.xml"))
        rels_xml = ElementTree.fromstring(archive.read("word/_rels/document.xml.rels"))
        package_rels_xml = ElementTree.fromstring(archive.read("_rels/.rels"))

    body = document_xml.find(qn("w:body"))
    assert body is not None
    paragraph = StdElementTree.Element(qn("w:p"))
    visible_run = StdElementTree.Element(qn("w:r"))
    visible_text = StdElementTree.Element(qn("w:t"))
    visible_text.text = "visible-before-change"
    visible_run.append(visible_text)
    insertion = StdElementTree.Element(qn("w:ins"))
    insertion.set(qn("w:id"), "1")
    inserted_run = StdElementTree.Element(qn("w:r"))
    inserted_text = StdElementTree.Element(qn("w:t"))
    inserted_text.text = "tracked-insertion"
    inserted_run.append(inserted_text)
    insertion.append(inserted_run)
    paragraph.extend((visible_run, insertion))
    body.insert(max(len(body) - 1, 0), paragraph)

    external = StdElementTree.SubElement(
        rels_xml, f"{{{RELATIONSHIPS_NS}}}Relationship"
    )
    external.set("Id", "rIdM6External")
    external.set(
        "Type",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/attachedTemplate",
    )
    external.set("Target", "https://example.invalid/remote-template.dotx")
    external.set("TargetMode", "External")

    custom = StdElementTree.SubElement(
        package_rels_xml, f"{{{RELATIONSHIPS_NS}}}Relationship"
    )
    custom.set("Id", "rIdM6Custom")
    custom.set("Type", "urn:reca:m6-spike:custom-part")
    custom.set("Target", "customXml/item99.xml")

    _rewrite_zip(
        path,
        {
            "word/document.xml": StdElementTree.tostring(document_xml),
            "word/_rels/document.xml.rels": StdElementTree.tostring(rels_xml),
            "_rels/.rels": StdElementTree.tostring(package_rels_xml),
        },
        {"customXml/item99.xml": b"<m6-spike keep='true' />"},
    )


def _candidate_docx_preflight(path: Path) -> tuple[str, ...]:
    limitations: list[str] = []
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 2_000:
            raise ValueError("too many archive entries")
        names = {entry.filename for entry in entries}
        total_uncompressed = 0
        for entry in entries:
            member = PurePosixPath(entry.filename.replace("\\", "/"))
            if member.is_absolute() or ".." in member.parts:
                raise ValueError("unsafe archive path")
            if len(member.parts) > 20:
                raise ValueError("archive path is too deep")
            mode = entry.external_attr >> 16
            if stat.S_ISLNK(mode):
                raise ValueError("symbolic links are not allowed")
            total_uncompressed += entry.file_size
            if total_uncompressed > 120_000_000:
                raise ValueError("archive expansion exceeds limit")
            if entry.file_size and entry.compress_size == 0:
                raise ValueError("invalid compression metadata")
            if entry.compress_size and entry.file_size / entry.compress_size > 100:
                raise ValueError("archive compression ratio exceeds limit")
            if PurePosixPath(entry.filename).suffix.lower() in {
                ".zip",
                ".docx",
                ".docm",
            }:
                raise ValueError("nested archives are not allowed")
            if entry.filename.endswith(".xml") and entry.file_size > 10_000_000:
                raise ValueError("XML part exceeds limit")

        if "[Content_Types].xml" not in names or "word/document.xml" not in names:
            raise ValueError("required DOCX parts are missing")
        content_types = ElementTree.fromstring(archive.read("[Content_Types].xml"))
        main_types = {
            node.attrib.get("ContentType")
            for node in content_types.findall(f"{{{CONTENT_TYPES_NS}}}Override")
            if node.attrib.get("PartName") == "/word/document.xml"
        }
        if main_types != {DOCX_CONTENT_TYPE}:
            raise ValueError("package is not a macro-free DOCX")
        if "word/vbaProject.bin" in names:
            raise ValueError("macro payload is not allowed")

        for name in sorted(names):
            if not name.endswith(".rels"):
                continue
            relationships = ElementTree.fromstring(archive.read(name))
            for relationship in relationships.findall(
                f"{{{RELATIONSHIPS_NS}}}Relationship"
            ):
                if relationship.attrib.get("TargetMode") == "External":
                    raise ValueError("external relationships are not allowed")

        if any(name.startswith("customXml/") for name in names):
            limitations.append("UNKNOWN_PARTS_PRESENT")
    return tuple(limitations)


def test_python_docx_314_dependency_surface() -> None:
    assert docx.__version__ == "1.2.0"
    assert lxml.__version__ == "6.1.1"


def test_deterministic_citation_and_numeric_tokenizer_boundary() -> None:
    text = (
        "王明等（2024）与 Smith et al., 2025b 报告 N=312，p < .05，"
        "r = 0.42，β=-0.18；见图 2 and Figure 3."
    )
    citations = [match.groupdict() for match in AUTHOR_YEAR.finditer(text)]
    numbers = [match.groupdict() for match in NUMERIC_FACT.finditer(text)]
    figures = [match.group("number") for match in FIGURE_REFERENCE.finditer(text)]

    assert citations == [
        {"author": "王明", "year": "2024"},
        {"author": "Smith", "year": "2025b"},
    ]
    assert numbers == [
        {"kind": "N", "value": "312"},
        {"kind": "p", "value": ".05"},
        {"kind": "r", "value": "0.42"},
        {"kind": "β", "value": "-0.18"},
    ]
    assert figures == ["2", "3"]


def test_supported_surface_and_round_trip_preservation(tmp_path: Path) -> None:
    source = tmp_path / "m6-spike.docx"
    document = Document()
    document.add_heading("M6 heading", level=1)
    paragraph = document.add_paragraph("N = 312, p = 0.03, r = 0.42, beta = 0.18")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " CITATION Example2026 ")
    field_run = OxmlElement("w:r")
    field_text = OxmlElement("w:t")
    field_text.text = "(Example, 2026)"
    field_run.append(field_text)
    field.append(field_run)
    paragraph._p.append(field)
    document.add_paragraph("Numbered item", style="List Number")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "metric"
    table.cell(0, 1).text = "value"
    document.add_picture(BytesIO(TINY_PNG))
    document.save(source)
    _add_spike_ooxml(source)

    loaded = Document(source)
    visible_text = [paragraph.text for paragraph in loaded.paragraphs]
    assert "M6 heading" in visible_text
    assert any("N = 312" in text for text in visible_text)
    assert "visible-before-change" in visible_text
    assert all("tracked-insertion" not in text for text in visible_text)
    assert len(loaded.tables) == 1
    assert len(loaded.inline_shapes) == 1

    round_tripped = tmp_path / "m6-spike-roundtrip.docx"
    loaded.save(round_tripped)
    with zipfile.ZipFile(round_tripped) as archive:
        assert archive.read("customXml/item99.xml") == b"<m6-spike keep='true' />"
        assert b"remote-template.dotx" in archive.read("word/_rels/document.xml.rels")


def test_candidate_preflight_rejects_external_relationships(tmp_path: Path) -> None:
    source = tmp_path / "external.docx"
    Document().save(source)
    _add_spike_ooxml(source)
    with pytest.raises(ValueError, match="external relationships"):
        _candidate_docx_preflight(source)


def test_candidate_preflight_reports_unknown_parts_without_executing_them(
    tmp_path: Path,
) -> None:
    source = tmp_path / "unknown-part.docx"
    Document().save(source)
    with zipfile.ZipFile(source) as archive:
        package_rels_xml = ElementTree.fromstring(archive.read("_rels/.rels"))
    custom = StdElementTree.SubElement(
        package_rels_xml, f"{{{RELATIONSHIPS_NS}}}Relationship"
    )
    custom.set("Id", "rIdM6Custom")
    custom.set("Type", "urn:reca:m6-spike:custom-part")
    custom.set("Target", "customXml/item99.xml")
    _rewrite_zip(
        source,
        {"_rels/.rels": StdElementTree.tostring(package_rels_xml)},
        {"customXml/item99.xml": b"<m6-spike keep='true' />"},
    )
    assert _candidate_docx_preflight(source) == ("UNKNOWN_PARTS_PRESENT",)


@pytest.mark.parametrize(
    ("member_name", "external_attr", "message"),
    [
        ("../escape.xml", 0, "unsafe archive path"),
        ("word/link.xml", (stat.S_IFLNK | 0o777) << 16, "symbolic links"),
        ("a/b/c/d/e/f/g/h/i/j/k/l/m/n/o/p/q/r/s/t/u.xml", 0, "too deep"),
    ],
)
def test_candidate_preflight_rejects_unsafe_members(
    tmp_path: Path, member_name: str, external_attr: int, message: str
) -> None:
    source = tmp_path / "unsafe.docx"
    Document().save(source)
    with zipfile.ZipFile(source, "a") as archive:
        info = zipfile.ZipInfo(member_name)
        info.external_attr = external_attr
        archive.writestr(info, b"unsafe")
    with pytest.raises(ValueError, match=message):
        _candidate_docx_preflight(source)


def test_candidate_preflight_rejects_macro_enabled_main_part(tmp_path: Path) -> None:
    source = tmp_path / "macro.docx"
    Document().save(source)
    with zipfile.ZipFile(source) as archive:
        content_types = archive.read("[Content_Types].xml").replace(
            DOCX_CONTENT_TYPE.encode(),
            b"application/vnd.ms-word.document.macroEnabled.main+xml",
        )
    _rewrite_zip(source, {"[Content_Types].xml": content_types}, {})
    with pytest.raises(ValueError, match="macro-free DOCX"):
        _candidate_docx_preflight(source)
