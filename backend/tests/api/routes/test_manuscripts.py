from __future__ import annotations

import hashlib
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any, cast
from zipfile import ZipFile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app import crud
from app.adapters.storage import StorageError, StorageObjectExists
from app.api.routes import manuscripts as manuscript_routes
from app.artifacts import service as artifact_service
from app.models import (
    Artifact,
    ArtifactStatus,
    ArtifactType,
    AuditLog,
    AuditResult,
    AuditResultStatus,
    Claim,
    ClaimStatus,
    Job,
    Manuscript,
    ManuscriptCheckRun,
    ManuscriptCheckRunStatus,
    ManuscriptIssue,
    ManuscriptIssueStatus,
    ManuscriptStatus,
    ManuscriptTransformation,
    ManuscriptTransformationStatus,
    ManuscriptVersion,
    ManuscriptVersionStatus,
    ManuscriptVersionType,
    StorageProvider,
    UserCreate,
)
from app.workers import jobs as worker_jobs


class MemoryStorage:
    def __init__(self) -> None:
        self.objects: dict[str, bytes] = {}

    def put_file_once(
        self, *, object_key: str, path: Path, content_sha256: str, size_bytes: int
    ) -> None:
        if object_key in self.objects:
            raise StorageObjectExists("exists")
        content = path.read_bytes()
        assert hashlib.sha256(content).hexdigest() == content_sha256
        assert len(content) == size_bytes
        self.objects[object_key] = content

    def download_to_path(self, *, object_key: str, path: Path) -> None:
        if object_key not in self.objects:
            raise StorageError("missing")
        path.write_bytes(self.objects[object_key])

    def presign_download(self, *, object_key: str, expires_seconds: int) -> str:
        return f"https://example.test/{object_key}?expires={expires_seconds}"


class RecordingDispatcher:
    def __init__(self) -> None:
        self.calls: list[tuple[uuid.UUID, str]] = []

    def dispatch(self, *, job_id: uuid.UUID, task_id: str) -> None:
        self.calls.append((job_id, task_id))


@pytest.fixture
def manuscript_runtime(
    monkeypatch: pytest.MonkeyPatch,
) -> tuple[MemoryStorage, RecordingDispatcher]:
    storage = MemoryStorage()
    dispatcher = RecordingDispatcher()
    monkeypatch.setattr(artifact_service, "storage", storage)
    monkeypatch.setattr(manuscript_routes, "dispatcher", dispatcher)
    return storage, dispatcher


def _docx() -> bytes:
    stream = BytesIO()
    document = Document()
    document.add_paragraph(
        "The intervention causes improvement in all adults (Smith, 2020).  N = 42."
    )
    document.add_heading("References", level=1)
    document.add_paragraph("Jones, A. (2021). Unused. doi:bad")
    document.save(stream)
    return stream.getvalue()


def _project(client: TestClient, headers: dict[str, str]) -> dict[str, Any]:
    response = client.post(
        "/api/v1/projects",
        headers={**headers, "Idempotency-Key": str(uuid.uuid4())},
        json={"name": "M6 manuscript", "project_type": "RESEARCH"},
    )
    assert response.status_code == 201, response.text
    return cast(dict[str, Any], response.json()["data"])


def _artifact(
    db: Session,
    storage: MemoryStorage,
    project_id: uuid.UUID,
    content: bytes,
    *,
    is_original: bool = True,
    source_artifact_id: uuid.UUID | None = None,
) -> Artifact:
    digest = hashlib.sha256(content).hexdigest()
    artifact = Artifact(
        project_id=project_id,
        artifact_type=ArtifactType.MANUSCRIPT_DOCX,
        filename="paper.docx",
        original_filename="paper.docx",
        storage_provider=StorageProvider.MINIO,
        storage_key=f"projects/{project_id}/m6/{uuid.uuid4()}",
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        size_bytes=len(content),
        sha256=digest,
        source_artifact_id=source_artifact_id,
        is_original=is_original,
        is_immutable=True,
        status=ArtifactStatus.AVAILABLE,
    )
    db.add(artifact)
    db.commit()
    storage.objects[artifact.storage_key] = content
    return artifact


def test_manuscript_check_vertical_chain_and_issue_decision(
    client: TestClient,
    db: Session,
    normal_user_token_headers: dict[str, str],
    manuscript_runtime: tuple[MemoryStorage, RecordingDispatcher],
) -> None:
    storage, dispatcher = manuscript_runtime
    project = _project(client, normal_user_token_headers)
    project_id = uuid.UUID(project["id"])
    artifact = _artifact(db, storage, project_id, _docx())
    created = client.post(
        f"/api/v1/projects/{project_id}/manuscripts",
        headers=normal_user_token_headers,
        json={"artifact_id": str(artifact.id), "title": "Fixture"},
    )
    assert created.status_code == 201, created.text
    version_id = created.json()["data"]["version"]["id"]
    assert created.json()["data"]["version"]["source_hash"] == artifact.sha256

    key = str(uuid.uuid4())
    request = {
        "checks": ["CITATION", "CAUSALITY", "BASIC_FORMAT"],
        "use_project_literature": True,
        "use_project_analysis_results": True,
        "use_project_figures": True,
    }
    accepted = client.post(
        f"/api/v1/manuscript-versions/{version_id}/check-runs",
        headers={**normal_user_token_headers, "Idempotency-Key": key},
        json=request,
    )
    replay = client.post(
        f"/api/v1/manuscript-versions/{version_id}/check-runs",
        headers={**normal_user_token_headers, "Idempotency-Key": key},
        json=request,
    )
    assert accepted.status_code == replay.status_code == 202
    assert replay.json()["meta"]["idempotency_replayed"] is True
    assert len(dispatcher.calls) == 1
    run_id = uuid.UUID(accepted.json()["data"]["manuscript_check_run"]["id"])
    job_id = uuid.UUID(accepted.json()["data"]["job"]["id"])
    assert worker_jobs._execute_job(object(), str(job_id))["completed"] is True
    db.expire_all()
    run = db.get(ManuscriptCheckRun, run_id)
    assert run is not None and run.status == ManuscriptCheckRunStatus.NEEDS_REVIEW
    issues = db.exec(
        select(ManuscriptIssue).where(ManuscriptIssue.manuscript_check_run_id == run_id)
    ).all()
    assert run.issue_count == len(issues) >= 4
    high = next(item for item in issues if item.severity.value == "HIGH")
    assert high.auto_fixable is False
    decision = client.post(
        f"/api/v1/manuscript-issues/{high.id}/accept",
        headers={**normal_user_token_headers, "If-Match": "1"},
        json={"reason": "Reviewed manually"},
    )
    assert decision.status_code == 200, decision.text
    assert decision.json()["data"]["status"] == ManuscriptIssueStatus.ACCEPTED
    assert db.get(ManuscriptVersion, uuid.UUID(version_id)).artifact_id == artifact.id
    assert (
        db.exec(select(AuditLog).where(AuditLog.object_id == high.id)).first()
        is not None
    )


def test_manuscript_worker_fails_closed_on_hash_change(
    client: TestClient,
    db: Session,
    normal_user_token_headers: dict[str, str],
    manuscript_runtime: tuple[MemoryStorage, RecordingDispatcher],
) -> None:
    storage, _ = manuscript_runtime
    project = _project(client, normal_user_token_headers)
    project_id = uuid.UUID(project["id"])
    artifact = _artifact(db, storage, project_id, _docx())
    created = client.post(
        f"/api/v1/projects/{project_id}/manuscripts",
        headers=normal_user_token_headers,
        json={"artifact_id": str(artifact.id)},
    )
    version_id = created.json()["data"]["version"]["id"]
    accepted = client.post(
        f"/api/v1/manuscript-versions/{version_id}/check-runs",
        headers={**normal_user_token_headers, "Idempotency-Key": str(uuid.uuid4())},
        json={"checks": ["CITATION"]},
    )
    run_id = uuid.UUID(accepted.json()["data"]["manuscript_check_run"]["id"])
    job_id = uuid.UUID(accepted.json()["data"]["job"]["id"])
    storage.objects[artifact.storage_key] = b"x" * artifact.size_bytes
    with pytest.raises(StorageError):
        worker_jobs._execute_job(object(), str(job_id))
    db.expire_all()
    assert db.get(ManuscriptCheckRun, run_id).status == ManuscriptCheckRunStatus.FAILED
    assert db.get(Job, job_id).status.value == "FAILED"


def test_project_manuscript_discovery_covers_empty_current_history_and_inactive(
    client: TestClient,
    db: Session,
    normal_user_token_headers: dict[str, str],
    manuscript_runtime: tuple[MemoryStorage, RecordingDispatcher],
) -> None:
    storage, _ = manuscript_runtime
    project = _project(client, normal_user_token_headers)
    project_id = uuid.UUID(project["id"])

    empty = client.get(
        f"/api/v1/projects/{project_id}/manuscript",
        headers=normal_user_token_headers,
    )
    assert empty.status_code == 200, empty.text
    assert empty.json()["data"] == {
        "state": "NONE",
        "current_manuscript": None,
        "current_version": None,
        "manuscripts": [],
        "versions": [],
    }

    artifact = _artifact(db, storage, project_id, _docx())
    created = client.post(
        f"/api/v1/projects/{project_id}/manuscripts",
        headers=normal_user_token_headers,
        json={"artifact_id": str(artifact.id), "title": "Discovery fixture"},
    )
    assert created.status_code == 201, created.text
    manuscript_id = uuid.UUID(created.json()["data"]["manuscript"]["id"])
    first_version_id = uuid.UUID(created.json()["data"]["version"]["id"])

    second_artifact = _artifact(
        db,
        storage,
        project_id,
        _docx(),
        is_original=False,
        source_artifact_id=artifact.id,
    )
    second_version = ManuscriptVersion(
        manuscript_id=manuscript_id,
        project_id=project_id,
        version_number=2,
        parent_version_id=first_version_id,
        artifact_id=second_artifact.id,
        version_type=ManuscriptVersionType.USER_REVISED,
        status=ManuscriptVersionStatus.AVAILABLE,
        source_hash=second_artifact.sha256,
    )
    db.add(second_version)
    db.flush()
    manuscript = db.get(Manuscript, manuscript_id)
    assert manuscript is not None
    manuscript.current_version_id = second_version.id
    db.add(manuscript)
    db.commit()

    current = client.get(
        f"/api/v1/projects/{project_id}/manuscript",
        headers=normal_user_token_headers,
    )
    assert current.status_code == 200, current.text
    data = current.json()["data"]
    assert data["state"] == "ACTIVE"
    assert data["current_manuscript"]["id"] == str(manuscript_id)
    assert data["current_version"]["id"] == str(second_version.id)
    assert [item["version_number"] for item in data["versions"]] == [1, 2]

    manuscript.status = ManuscriptStatus.ARCHIVED
    db.add(manuscript)
    db.commit()
    archived = client.get(
        f"/api/v1/projects/{project_id}/manuscript",
        headers=normal_user_token_headers,
    ).json()["data"]
    assert archived["state"] == "ARCHIVED"
    assert archived["current_manuscript"] is None
    assert archived["current_version"] is None
    assert archived["manuscripts"][0]["status"] == "ARCHIVED"
    assert len(archived["versions"]) == 2

    manuscript.status = ManuscriptStatus.INVALIDATED
    db.add(manuscript)
    db.commit()
    invalidated = client.get(
        f"/api/v1/projects/{project_id}/manuscript",
        headers=normal_user_token_headers,
    ).json()["data"]
    assert invalidated["state"] == "INVALIDATED"
    assert invalidated["current_manuscript"] is None
    assert invalidated["manuscripts"][0]["status"] == "INVALIDATED"


def test_manuscript_no_disclosure_for_non_member(
    client: TestClient,
    db: Session,
    normal_user_token_headers: dict[str, str],
    manuscript_runtime: tuple[MemoryStorage, RecordingDispatcher],
) -> None:
    storage, _ = manuscript_runtime
    project = _project(client, normal_user_token_headers)
    artifact = _artifact(db, storage, uuid.UUID(project["id"]), _docx())
    created = client.post(
        f"/api/v1/projects/{project['id']}/manuscripts",
        headers=normal_user_token_headers,
        json={"artifact_id": str(artifact.id)},
    )
    manuscript_id = created.json()["data"]["manuscript"]["id"]
    version_id = created.json()["data"]["version"]["id"]
    claim = client.post(
        f"/api/v1/projects/{project['id']}/claims",
        headers={
            **normal_user_token_headers,
            "Idempotency-Key": str(uuid.uuid4()),
        },
        json={
            "claim_type": "MANUSCRIPT_STATEMENT",
            "source_object_type": "manuscript_version",
            "source_object_id": version_id,
            "source_location": {"paragraph_index": 0},
            "claim_text": "The intervention causes improvement",
        },
    )
    assert claim.status_code == 201, claim.text
    outsider_email = f"outsider-{uuid.uuid4()}@example.com"
    outsider = crud.create_user(
        session=db,
        user_create=UserCreate(email=outsider_email, password="outsider-password"),
    )
    del outsider
    from tests.utils.user import user_authentication_headers

    outsider_headers = user_authentication_headers(
        client=client, email=outsider_email, password="outsider-password"
    )
    assert (
        client.get(
            f"/api/v1/manuscripts/{manuscript_id}", headers=outsider_headers
        ).status_code
        == 404
    )
    assert (
        client.get(
            f"/api/v1/projects/{project['id']}/manuscript",
            headers=outsider_headers,
        ).status_code
        == 404
    )
    assert (
        client.get(
            f"/api/v1/claims/{claim.json()['data']['id']}", headers=outsider_headers
        ).status_code
        == 404
    )


def test_revision_audit_and_claim_confirmation_backend_closure(
    client: TestClient,
    db: Session,
    normal_user_token_headers: dict[str, str],
    manuscript_runtime: tuple[MemoryStorage, RecordingDispatcher],
) -> None:
    storage, dispatcher = manuscript_runtime
    project = _project(client, normal_user_token_headers)
    project_id = uuid.UUID(project["id"])
    before_artifact = _artifact(db, storage, project_id, _docx())
    created = client.post(
        f"/api/v1/projects/{project_id}/manuscripts",
        headers=normal_user_token_headers,
        json={"artifact_id": str(before_artifact.id)},
    )
    assert created.status_code == 201, created.text
    manuscript_id = uuid.UUID(created.json()["data"]["manuscript"]["id"])
    before_id = uuid.UUID(created.json()["data"]["version"]["id"])

    revised = BytesIO()
    revised_doc = Document()
    revised_doc.add_paragraph(
        "Everyone is affected by the intervention (Jones, 2021).  N = 40. Figure 2."
    )
    revised_doc.save(revised)
    revised_bytes = revised.getvalue()
    after_artifact = _artifact(
        db,
        storage,
        project_id,
        revised_bytes,
        is_original=False,
        source_artifact_id=before_artifact.id,
    )
    after = ManuscriptVersion(
        manuscript_id=manuscript_id,
        project_id=project_id,
        version_number=2,
        parent_version_id=before_id,
        artifact_id=after_artifact.id,
        version_type=ManuscriptVersionType.USER_REVISED,
        status=ManuscriptVersionStatus.AVAILABLE,
        source_hash=after_artifact.sha256,
    )
    db.add(after)
    db.commit()

    accepted = client.post(
        f"/api/v1/projects/{project_id}/manuscript-revision-audits",
        headers={
            **normal_user_token_headers,
            "Idempotency-Key": str(uuid.uuid4()),
        },
        json={
            "baseline_manuscript_version_id": str(before_id),
            "candidate_manuscript_version_id": str(after.id),
            "referenced_result_ids": [],
        },
    )
    assert accepted.status_code == 202, accepted.text
    assert dispatcher.calls
    audit_id = uuid.UUID(accepted.json()["data"]["audit_result"]["id"])
    audit_job_id = uuid.UUID(accepted.json()["data"]["job"]["id"])
    assert worker_jobs._execute_job(object(), str(audit_job_id))["completed"] is True
    db.expire_all()
    audit = db.get(AuditResult, audit_id)
    assert audit is not None and audit.status == AuditResultStatus.COMPLETED
    assert {item["code"] for item in audit.result["findings"]} >= {
        "CLAIM_NUMERIC_MISMATCH",
        "CITATION_SET_CHANGED",
    }
    assert db.get(Manuscript, manuscript_id).current_version_id == before_id

    claim_created = client.post(
        f"/api/v1/projects/{project_id}/claims",
        headers={
            **normal_user_token_headers,
            "Idempotency-Key": str(uuid.uuid4()),
        },
        json={
            "claim_type": "MANUSCRIPT_STATEMENT",
            "source_object_type": "manuscript_version",
            "source_object_id": str(before_id),
            "source_location": {"paragraph_index": 0},
            "claim_text": "The intervention causes improvement",
            "scope_statement": "Current manuscript wording.",
        },
    )
    assert claim_created.status_code == 201, claim_created.text
    claim_id = uuid.UUID(claim_created.json()["data"]["id"])
    supported = client.patch(
        f"/api/v1/claims/{claim_id}",
        headers={**normal_user_token_headers, "If-Match": "1"},
        json={"status": "SUPPORTED", "confidence": "HIGH"},
    )
    assert supported.status_code == 200, supported.text
    stale_patch = client.patch(
        f"/api/v1/claims/{claim_id}",
        headers={**normal_user_token_headers, "If-Match": "1"},
        json={"scope_statement": "Stale write"},
    )
    assert stale_patch.status_code == 412
    direct_confirmation = client.patch(
        f"/api/v1/claims/{claim_id}",
        headers={**normal_user_token_headers, "If-Match": "2"},
        json={"status": "CONFIRMED"},
    )
    assert direct_confirmation.status_code == 409
    assert direct_confirmation.json()["error"]["code"] == "APPROVAL_REQUIRED"
    requested = client.post(
        f"/api/v1/claims/{claim_id}/confirmation-requests",
        headers={
            **normal_user_token_headers,
            "Idempotency-Key": str(uuid.uuid4()),
        },
    )
    assert requested.status_code == 201, requested.text
    approval_id = requested.json()["data"]["approval_id"]
    approved = client.post(
        f"/api/v1/approvals/{approval_id}/approve",
        headers={
            **normal_user_token_headers,
            "Idempotency-Key": str(uuid.uuid4()),
        },
        json={"decision_reason": "Source checked", "item_decisions": []},
    )
    assert approved.status_code == 200, approved.text
    db.expire_all()
    claim = db.get(Claim, claim_id)
    assert claim is not None and claim.status == ClaimStatus.CONFIRMED
    terminal_edit = client.patch(
        f"/api/v1/claims/{claim_id}",
        headers={
            **normal_user_token_headers,
            "If-Match": str(claim.lock_version),
        },
        json={"scope_statement": "Must not edit"},
    )
    assert terminal_edit.status_code == 409


def test_low_risk_fix_preview_approval_and_atomic_version(
    client: TestClient,
    db: Session,
    normal_user_token_headers: dict[str, str],
    manuscript_runtime: tuple[MemoryStorage, RecordingDispatcher],
) -> None:
    storage, _ = manuscript_runtime
    project = _project(client, normal_user_token_headers)
    project_id = uuid.UUID(project["id"])
    artifact = _artifact(db, storage, project_id, _docx())
    created = client.post(
        f"/api/v1/projects/{project_id}/manuscripts",
        headers=normal_user_token_headers,
        json={"artifact_id": str(artifact.id)},
    )
    version_id = uuid.UUID(created.json()["data"]["version"]["id"])
    checked = client.post(
        f"/api/v1/manuscript-versions/{version_id}/check-runs",
        headers={
            **normal_user_token_headers,
            "Idempotency-Key": str(uuid.uuid4()),
        },
        json={"checks": ["BASIC_FORMAT"]},
    )
    check_job_id = uuid.UUID(checked.json()["data"]["job"]["id"])
    assert worker_jobs._execute_job(object(), str(check_job_id))["completed"] is True
    run_id = uuid.UUID(checked.json()["data"]["manuscript_check_run"]["id"])
    db.expire_all()
    issue = db.exec(
        select(ManuscriptIssue).where(
            ManuscriptIssue.manuscript_check_run_id == run_id,
            ManuscriptIssue.auto_fixable.is_(True),
        )
    ).one()
    accepted = client.post(
        f"/api/v1/manuscript-issues/{issue.id}/accept",
        headers={**normal_user_token_headers, "If-Match": "1"},
        json={"reason": "Low-risk formatting"},
    )
    assert accepted.status_code == 200, accepted.text
    plan_response = client.post(
        f"/api/v1/manuscript-versions/{version_id}/fix-plans",
        headers={
            **normal_user_token_headers,
            "Idempotency-Key": str(uuid.uuid4()),
        },
        json={"issue_ids": [str(issue.id)]},
    )
    assert plan_response.status_code == 201, plan_response.text
    plan_id = uuid.UUID(plan_response.json()["data"]["id"])
    preview = client.post(
        f"/api/v1/manuscript-fix-plans/{plan_id}/preview",
        headers={**normal_user_token_headers, "If-Match": "1"},
    )
    assert preview.status_code == 200, preview.text
    assert len(storage.objects) == 1
    requested = client.post(
        f"/api/v1/manuscript-fix-plans/{plan_id}/approval-requests",
        headers={
            **normal_user_token_headers,
            "Idempotency-Key": str(uuid.uuid4()),
        },
    )
    assert requested.status_code == 201, requested.text
    approval_id = requested.json()["data"]["approval_id"]
    approved = client.post(
        f"/api/v1/approvals/{approval_id}/approve",
        headers={
            **normal_user_token_headers,
            "Idempotency-Key": str(uuid.uuid4()),
        },
        json={
            "decision_reason": "Formatting only",
            "item_decisions": [
                {
                    "item_type": "manuscript_issue",
                    "item_id": str(issue.id),
                    "decision": "ACCEPT",
                }
            ],
        },
    )
    assert approved.status_code == 200, approved.text
    execute = client.post(
        f"/api/v1/manuscript-fix-plans/{plan_id}/execute",
        headers={
            **normal_user_token_headers,
            "Idempotency-Key": str(uuid.uuid4()),
        },
    )
    assert execute.status_code == 202, execute.text
    job_id = uuid.UUID(execute.json()["data"]["job"]["id"])
    assert worker_jobs._execute_job(object(), str(job_id))["completed"] is True
    db.expire_all()
    plan = db.get(ManuscriptTransformation, plan_id)
    assert plan is not None and plan.status == ManuscriptTransformationStatus.COMPLETED
    output = db.get(ManuscriptVersion, plan.output_manuscript_version_id)
    assert (
        output is not None and output.version_type == ManuscriptVersionType.AUTO_FIXED
    )
    assert output.parent_version_id == version_id
    assert db.get(Manuscript, output.manuscript_id).current_version_id == output.id
    assert db.get(ManuscriptIssue, issue.id).status == ManuscriptIssueStatus.RESOLVED
    assert len(storage.objects) == 2
    output_artifact = db.get(Artifact, output.artifact_id)
    assert output_artifact is not None
    with ZipFile(BytesIO(storage.objects[artifact.storage_key])) as source_package:
        with ZipFile(
            BytesIO(storage.objects[output_artifact.storage_key])
        ) as output_package:
            custom_parts = [
                name
                for name in source_package.namelist()
                if name.startswith("customXml/")
            ]
            assert custom_parts
            assert {name: source_package.read(name) for name in custom_parts} == {
                name: output_package.read(name) for name in custom_parts
            }
