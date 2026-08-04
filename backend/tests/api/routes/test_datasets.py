from __future__ import annotations

import hashlib
import io
import uuid
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlmodel import Session, select

from app import crud
from app.adapters.storage import StorageError, StorageObjectExists
from app.artifacts import service as artifact_service
from app.models import (
    Artifact,
    AuditLog,
    Dataset,
    DatasetColumn,
    DatasetVersion,
    UserCreate,
)
from tests.utils.user import user_authentication_headers
from tests.utils.utils import random_lower_string


class MemoryStorage:
    def __init__(self) -> None:
        self.objects: dict[str, bytes] = {}

    def put_file_once(
        self, *, object_key: str, path: Path, content_sha256: str, size_bytes: int
    ) -> None:
        content = path.read_bytes()
        if object_key in self.objects:
            raise StorageObjectExists("already exists")
        assert len(content) == size_bytes
        assert hashlib.sha256(content).hexdigest() == content_sha256
        self.objects[object_key] = content

    def download_to_path(self, *, object_key: str, path: Path) -> None:
        try:
            path.write_bytes(self.objects[object_key])
        except KeyError as exc:
            raise StorageError("missing") from exc

    def presign_download(self, *, object_key: str, expires_seconds: int) -> str:
        return f"https://example.test/{object_key}?expires={expires_seconds}"


@pytest.fixture
def memory_storage(monkeypatch: pytest.MonkeyPatch) -> MemoryStorage:
    backend = MemoryStorage()
    monkeypatch.setattr(artifact_service, "storage", backend)
    return backend


def _project(client: TestClient, headers: dict[str, str]) -> dict[str, object]:
    response = client.post(
        "/api/v1/projects",
        headers={**headers, "Idempotency-Key": str(uuid.uuid4())},
        json={"name": "Dataset project", "project_type": "RESEARCH"},
    )
    assert response.status_code == 201, response.text
    return response.json()["data"]


def _upload(
    client: TestClient,
    headers: dict[str, str],
    project_id: object,
    content: bytes,
    *,
    filename: str = "data.csv",
    content_type: str = "text/csv",
    key: str | None = None,
) -> object:
    return client.post(
        f"/api/v1/projects/{project_id}/datasets",
        headers={**headers, "Idempotency-Key": key or str(uuid.uuid4())},
        files={"file": (filename, content, content_type)},
        data={
            "name": "Survey",
            "source_type": "USER_UPLOAD",
            "license_status": "UNKNOWN",
        },
    )


def _xlsx() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    first.append(["id", "score"])
    first.append([1, 10])
    second = workbook.create_sheet("Second")
    second.append(["id", "label"])
    second.append([2, "B"])
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["secret"])
    hidden.append(["value"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_csv_upload_preview_columns_etag_and_idempotency(
    client: TestClient,
    db: Session,
    normal_user_token_headers: dict[str, str],
    memory_storage: MemoryStorage,
) -> None:
    assert artifact_service.storage is memory_storage
    project = _project(client, normal_user_token_headers)
    content = b"id,name,score\n1,Alice,10\n2,Bob,\n"
    key = str(uuid.uuid4())
    first = _upload(client, normal_user_token_headers, project["id"], content, key=key)
    replay = _upload(client, normal_user_token_headers, project["id"], content, key=key)
    assert first.status_code == replay.status_code == 201, first.text
    assert replay.json()["meta"]["idempotency_replayed"] is True
    data = first.json()["data"]
    version = data["version"]
    assert version["status"] == "AVAILABLE"
    assert version["version_type"] == "ORIGINAL"
    assert version["parent_version_id"] is None
    assert version["data_hash"] == hashlib.sha256(content).hexdigest()
    dataset = db.get(Dataset, uuid.UUID(data["dataset"]["id"]))
    artifact = db.get(Artifact, uuid.UUID(version["artifact_id"]))
    assert dataset is not None and artifact is not None
    assert dataset.current_version_id == uuid.UUID(version["id"])
    assert artifact.sha256 == version["data_hash"]
    assert artifact.is_original is artifact.is_immutable is True
    assert memory_storage.objects[artifact.storage_key] == content

    preview = client.get(
        f"/api/v1/dataset-versions/{version['id']}/preview",
        headers=normal_user_token_headers,
        params={"limit": 1, "columns": ["id", "score"]},
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["data"]["rows"] == [{"id": "1", "score": "10"}]

    columns = client.get(
        f"/api/v1/dataset-versions/{version['id']}/columns",
        headers=normal_user_token_headers,
    )
    assert columns.status_code == 200
    score = next(
        item for item in columns.json()["data"] if item["source_name"] == "score"
    )
    assert score["inferred_type"] == "INTEGER"
    assert score["confirmed_type"] is None
    update = client.patch(
        f"/api/v1/dataset-columns/{score['id']}",
        headers={**normal_user_token_headers, "If-Match": '"1"'},
        json={
            "confirmed_type": "NUMERIC",
            "semantic_role": "DEPENDENT_VARIABLE",
            "confirmation_status": "CONFIRMED",
        },
    )
    assert update.status_code == 200, update.text
    assert update.json()["data"]["inferred_type"] == "INTEGER"
    assert update.json()["data"]["confirmed_type"] == "NUMERIC"
    stale = client.patch(
        f"/api/v1/dataset-columns/{score['id']}",
        headers={**normal_user_token_headers, "If-Match": '"1"'},
        json={"display_name": "Score"},
    )
    assert stale.status_code == 409
    assert stale.json()["error"]["code"] == "RESOURCE_VERSION_CONFLICT"

    invalidation_key = str(uuid.uuid4())
    invalidated = client.post(
        f"/api/v1/dataset-versions/{version['id']}/invalidate",
        headers={**normal_user_token_headers, "Idempotency-Key": invalidation_key},
        json={"reason": "Source correction required"},
    )
    invalidation_replay = client.post(
        f"/api/v1/dataset-versions/{version['id']}/invalidate",
        headers={**normal_user_token_headers, "Idempotency-Key": invalidation_key},
        json={"reason": "Source correction required"},
    )
    assert invalidated.status_code == invalidation_replay.status_code == 200
    assert invalidated.json()["data"]["status"] == "INVALIDATED"
    assert invalidation_replay.json()["meta"]["idempotency_replayed"] is True
    db.refresh(dataset)
    assert dataset.current_version_id is None
    assert db.exec(
        select(AuditLog).where(
            AuditLog.object_id == uuid.UUID(version["id"]),
            AuditLog.action == "DATASET_VERSION_INVALIDATED",
        )
    ).one()


def test_multisheet_selection_is_persisted_and_hidden_requires_ack(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    memory_storage: MemoryStorage,
) -> None:
    assert artifact_service.storage is memory_storage
    project = _project(client, normal_user_token_headers)
    uploaded = _upload(
        client,
        normal_user_token_headers,
        project["id"],
        _xlsx(),
        filename="book.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    assert uploaded.status_code == 201, uploaded.text
    version = uploaded.json()["data"]["version"]
    assert version["status"] == "CREATING"
    assert version["selected_worksheet_name"] is None
    worksheets = client.get(
        f"/api/v1/dataset-versions/{version['id']}/worksheets",
        headers=normal_user_token_headers,
    )
    assert [item["visibility"] for item in worksheets.json()["data"]["worksheets"]] == [
        "VISIBLE",
        "VISIBLE",
        "HIDDEN",
    ]

    hidden_without_ack = client.post(
        f"/api/v1/dataset-versions/{version['id']}/worksheet-selection",
        headers={**normal_user_token_headers, "Idempotency-Key": str(uuid.uuid4())},
        json={"worksheet_name": "Hidden", "acknowledge_hidden": False},
    )
    assert hidden_without_ack.status_code == 422
    assert hidden_without_ack.json()["error"]["code"] == "HIDDEN_WORKSHEET_ACK_REQUIRED"

    selected = client.post(
        f"/api/v1/dataset-versions/{version['id']}/worksheet-selection",
        headers={**normal_user_token_headers, "Idempotency-Key": str(uuid.uuid4())},
        json={"worksheet_name": "Second", "acknowledge_hidden": False},
    )
    assert selected.status_code == 201, selected.text
    assert selected.json()["data"]["status"] == "AVAILABLE"
    assert selected.json()["data"]["selected_worksheet_name"] == "Second"


def test_dataset_no_disclosure_and_permissions(
    client: TestClient,
    db: Session,
    normal_user_token_headers: dict[str, str],
    memory_storage: MemoryStorage,
) -> None:
    assert artifact_service.storage is memory_storage
    project = _project(client, normal_user_token_headers)
    uploaded = _upload(client, normal_user_token_headers, project["id"], b"id\n1\n")
    dataset_id = uploaded.json()["data"]["dataset"]["id"]
    version_id = uploaded.json()["data"]["version"]["id"]
    password = random_lower_string()
    outsider = crud.create_user(
        session=db,
        user_create=UserCreate(
            email=f"{random_lower_string()}@example.com", password=password
        ),
    )
    outsider_headers = user_authentication_headers(
        client=client, email=outsider.email, password=password
    )
    assert (
        client.get(
            f"/api/v1/datasets/{dataset_id}", headers=outsider_headers
        ).status_code
        == 404
    )
    assert (
        client.get(
            f"/api/v1/datasets/{dataset_id}/versions", headers=outsider_headers
        ).status_code
        == 404
    )
    assert (
        client.patch(
            f"/api/v1/datasets/{dataset_id}",
            headers={**outsider_headers, "If-Match": '"1"'},
            json={"name": "No"},
        ).status_code
        == 404
    )

    viewer_password = random_lower_string()
    viewer = crud.create_user(
        session=db,
        user_create=UserCreate(
            email=f"{random_lower_string()}@example.com", password=viewer_password
        ),
    )
    added = client.post(
        f"/api/v1/projects/{project['id']}/members",
        headers={**normal_user_token_headers, "Idempotency-Key": str(uuid.uuid4())},
        json={"user_id": str(viewer.id), "role": "VIEWER"},
    )
    assert added.status_code == 201
    viewer_headers = user_authentication_headers(
        client=client, email=viewer.email, password=viewer_password
    )
    assert (
        client.get(f"/api/v1/datasets/{dataset_id}", headers=viewer_headers).status_code
        == 200
    )
    history = client.get(
        f"/api/v1/datasets/{dataset_id}/versions", headers=viewer_headers
    )
    assert history.status_code == 200
    assert history.json()["meta"]["count"] == 1
    assert history.json()["data"][0]["id"] == version_id
    forbidden_upload = _upload(client, viewer_headers, project["id"], b"id\n2\n")
    assert forbidden_upload.status_code == 403
    forbidden_update = client.patch(
        f"/api/v1/datasets/{dataset_id}",
        headers={**viewer_headers, "If-Match": '"1"'},
        json={"name": "No"},
    )
    assert forbidden_update.status_code == 403

    missing = client.patch(
        f"/api/v1/datasets/{dataset_id}",
        headers=normal_user_token_headers,
        json={"name": "New"},
    )
    assert missing.status_code == 400
    updated = client.patch(
        f"/api/v1/datasets/{dataset_id}",
        headers={**normal_user_token_headers, "If-Match": '"1"'},
        json={"name": "New"},
    )
    assert updated.status_code == 200
    assert updated.headers["etag"] == '"2"'
    assert updated.json()["data"]["license_status"] == "UNKNOWN"


def test_original_model_constraints_are_exposed(db: Session) -> None:
    assert DatasetVersion.__table__.constraints
    assert DatasetColumn.__table__.constraints
    assert db.exec(select(DatasetVersion)).all() is not None


def test_upload_byte_limit_and_mime_mismatch_fail_explicitly(
    client: TestClient,
    normal_user_token_headers: dict[str, str],
    memory_storage: MemoryStorage,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    assert artifact_service.storage is memory_storage
    project = _project(client, normal_user_token_headers)
    monkeypatch.setattr("app.datasets.limits.MAX_UPLOAD_BYTES", 3)
    too_large = _upload(client, normal_user_token_headers, project["id"], b"id\n1\n")
    assert too_large.status_code == 413
    monkeypatch.setattr("app.datasets.limits.MAX_UPLOAD_BYTES", 25 * 1024 * 1024)
    mismatch = _upload(
        client,
        normal_user_token_headers,
        project["id"],
        b"id\n1\n",
        content_type="image/png",
    )
    assert mismatch.status_code == 422
    assert mismatch.json()["error"]["code"] == "FILE_TYPE_MISMATCH"
