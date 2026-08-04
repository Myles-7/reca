from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.api.errors import ContractError
from app.datasets import parsers
from app.models import DatasetColumnType

pytestmark = pytest.mark.no_database


def test_csv_bom_legacy_encoding_and_stable_profile(tmp_path: Path) -> None:
    bom = tmp_path / "bom.csv"
    bom.write_bytes("name,score\nAlice,1\nBob,\n".encode("utf-8-sig"))
    legacy = tmp_path / "legacy.csv"
    legacy.write_bytes("城市;值\n北京;3\n上海;4\n".encode("gb18030"))

    first = parsers.parse_csv(bom)
    second = parsers.parse_csv(bom)
    legacy_profile = parsers.parse_csv(legacy)

    assert first == second
    assert first.row_count == 2
    assert first.columns[1].inferred_type == DatasetColumnType.INTEGER
    assert first.columns[1].missing_ratio == 0.5
    assert legacy_profile.headers == ("城市", "值")


def test_xlsx_manifest_lists_hidden_sheets_and_rejects_formulas(tmp_path: Path) -> None:
    workbook = Workbook()
    visible = workbook.active
    visible.title = "Visible"
    visible.append(["id", "value"])
    visible.append([1, 2])
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["id"])
    hidden.append([1])
    path = tmp_path / "sheets.xlsx"
    workbook.save(path)

    manifest = parsers.worksheet_manifest(path)
    assert [(item.name, item.visibility) for item in manifest] == [
        ("Visible", "VISIBLE"),
        ("Hidden", "HIDDEN"),
    ]
    assert parsers.parse_xlsx(path, "Visible").row_count == 1

    visible["B2"] = "=1+1"
    formula = tmp_path / "formula.xlsx"
    workbook.save(formula)
    with pytest.raises(ContractError, match="formulas") as error:
        parsers.worksheet_manifest(formula)
    assert error.value.code == "UNSAFE_WORKBOOK"


def test_type_disguise_damage_and_limits_fail_closed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    disguised = tmp_path / "fake.xlsx"
    disguised.write_text("a,b\n1,2\n", encoding="utf-8")
    with pytest.raises(ContractError) as mismatch:
        parsers.detect_format(
            disguised.name,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            disguised.read_bytes()[:8],
        )
    assert mismatch.value.code == "FILE_TYPE_MISMATCH"

    damaged = tmp_path / "damaged.xlsx"
    damaged.write_bytes(b"PK\x03\x04not-a-workbook")
    with pytest.raises(ContractError) as invalid:
        parsers.worksheet_manifest(damaged)
    assert invalid.value.code == "INVALID_WORKBOOK"

    too_wide = tmp_path / "wide.csv"
    too_wide.write_text(
        ",".join(f"c{i}" for i in range(101))
        + "\n"
        + ",".join("1" for _ in range(101)),
        encoding="utf-8",
    )
    with pytest.raises(ContractError) as wide:
        parsers.parse_csv(too_wide)
    assert wide.value.code == "COLUMN_LIMIT_EXCEEDED"

    rows = tmp_path / "rows.csv"
    rows.write_text("id\n1\n2\n", encoding="utf-8")
    monkeypatch.setattr("app.datasets.limits.MAX_ROWS", 1)
    with pytest.raises(ContractError) as too_many_rows:
        parsers.parse_csv(rows)
    assert too_many_rows.value.code == "ROW_LIMIT_EXCEEDED"

    monkeypatch.setattr("app.datasets.limits.MAX_ROWS", 50_000)
    monkeypatch.setattr("app.datasets.limits.MAX_CELL_CHARS", 3)
    cell = tmp_path / "cell.csv"
    cell.write_text("value\nlong\n", encoding="utf-8")
    with pytest.raises(ContractError) as too_long:
        parsers.parse_csv(cell)
    assert too_long.value.code == "CELL_LIMIT_EXCEEDED"


def test_xlsx_external_links_and_compression_risk_are_rejected(
    tmp_path: Path,
) -> None:
    workbook = Workbook()
    workbook.active.append(["id"])
    workbook.active.append([1])
    external = tmp_path / "external.xlsx"
    workbook.save(external)
    with zipfile.ZipFile(external, "a") as archive:
        archive.writestr("xl/externalLinks/externalLink1.xml", "<externalLink/>")
    with pytest.raises(ContractError) as link_error:
        parsers.worksheet_manifest(external)
    assert link_error.value.code == "UNSAFE_WORKBOOK"

    compressed = tmp_path / "compressed.xlsx"
    with zipfile.ZipFile(compressed, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "x")
        archive.writestr("xl/worksheets/sheet1.xml", b"0" * 1_000_000)
    with pytest.raises(ContractError) as compression_error:
        parsers.worksheet_manifest(compressed)
    assert compression_error.value.code == "WORKBOOK_COMPRESSION_RISK"
